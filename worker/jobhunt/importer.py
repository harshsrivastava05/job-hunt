"""Load companies + careers URLs from CSV/XLSX. Safe to re-run; a changed URL moves the company.

Usage: python -m jobhunt.importer <companies.csv|companies.xlsx>
"""
import csv
import io
import sys
from collections import Counter
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

from jobhunt.db import connect

REQUIRED = {"company_name", "careers_url"}


def normalize_url(raw):
    raw = raw.strip()
    if "://" not in raw:
        raw = "https://" + raw
    p = urlsplit(raw)
    if p.scheme not in ("http", "https") or "." not in p.netloc or any(c.isspace() for c in p.netloc):
        return None
    return urlunsplit((p.scheme, p.netloc.lower(), p.path.rstrip("/"), p.query, p.fragment))


def read_rows(path):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        # not read_only: that mode keeps the file open (locked on Windows) until wb.close()
        rows = load_workbook(path, data_only=True).active.iter_rows(values_only=True)
        header = [str(h or "") for h in next(rows, [])]
        records = [dict(zip(header, ("" if v is None else str(v) for v in r))) for r in rows]
    elif path.suffix.lower() == ".csv":
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:  # Excel's "CSV" export on Windows
            text = path.read_text(encoding="cp1252")
        records = list(csv.DictReader(io.StringIO(text)))
    else:
        sys.exit(f"{path.name}: expected .csv or .xlsx (save old .xls files as .xlsx)")
    return [{k.strip().lower(): (v or "").strip() for k, v in r.items() if k} for r in records]


def import_file(conn, path):
    rows = read_rows(path)
    if rows and not REQUIRED <= rows[0].keys():
        sys.exit(f"{Path(path).name}: missing column(s) {sorted(REQUIRED - rows[0].keys())}")
    stats = Counter()
    with conn:  # all rows or none
        for r in rows:
            if not any(r.values()):  # trailing blank spreadsheet rows
                continue
            name, city, url = r["company_name"], r.get("city", ""), normalize_url(r["careers_url"])
            if not name or not url:
                stats["skipped"] += 1
                continue
            stats["new_sources"] += conn.execute("INSERT OR IGNORE INTO sources (url) VALUES (?)", (url,)).rowcount
            source_id = conn.execute("SELECT id FROM sources WHERE url = ?", (url,)).fetchone()[0]
            old = conn.execute("SELECT source_id FROM companies WHERE name = ? AND city = ?", (name, city)).fetchone()
            if old is None:
                conn.execute("INSERT INTO companies (name, city, source_id) VALUES (?, ?, ?)", (name, city, source_id))
                stats["new_companies"] += 1
            elif old[0] != source_id:
                conn.execute("UPDATE companies SET source_id = ? WHERE name = ? AND city = ?", (source_id, name, city))
                stats["moved"] += 1
            else:
                stats["unchanged"] += 1
    return stats


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    conn = connect()
    stats = import_file(conn, sys.argv[1])
    sources, live, companies = conn.execute(
        "SELECT count(*), sum(EXISTS (SELECT 1 FROM companies c WHERE c.source_id = s.id)),"
        " (SELECT count(*) FROM companies) FROM sources s"
    ).fetchone()
    print(dict(stats))
    print(f"db: {companies} companies, {sources} sources ({live} watched)")
